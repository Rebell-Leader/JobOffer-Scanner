"""Resume ingestion + ATS/keyword-gap analysis.

Pipeline:
  1. ``extract_resume_text`` parses an uploaded resume (PDF / DOCX / plain text)
     into a normalized string. Optional deps (``pypdf``, ``python-docx``) are
     imported lazily so the rest of the app keeps working without them.
  2. ``ats_keyword_match`` computes a deterministic match score between the
     resume text and the required skills extracted from the job posting —
     pure logic, fully unit-testable.
  3. ``ats_format_checks`` flags formatting patterns that historically trip
     ATS parsers (tables, multi-column hints, excessive ALL-CAPS lines).
  4. ``analyze_resume`` ties these together and asks the LLM for the prose
     "gap analysis" section, with the resume text wrapped as untrusted data.
"""

from __future__ import annotations

import io
import logging
import re
from typing import Dict, Iterable, List, Optional

from utils.cache import cache
from utils.llm import get_completion
from utils.security import wrap_untrusted

logger = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# 1. Text extraction
# ---------------------------------------------------------------------------

def extract_resume_text(file_bytes: bytes, filename: str) -> str:
    """Extract plain text from a resume upload.

    Supports .pdf, .docx, .txt, .md. Raises ``ValueError`` for unsupported
    types or unreadable files so the caller can show a clear error rather
    than silently analyzing an empty document.
    """
    name = (filename or "").lower()
    if name.endswith((".txt", ".md")):
        return file_bytes.decode("utf-8", errors="replace")
    if not name.endswith((".pdf", ".docx")):
        raise ValueError(f"Unsupported resume format: {filename!r}. Use PDF, DOCX, or TXT.")
    try:
        text = _extract_pdf(file_bytes) if name.endswith(".pdf") else _extract_docx(file_bytes)
    except ValueError:
        raise
    except Exception as exc:  # noqa: BLE001 - corrupt/encrypted files raise library-specific errors
        logger.warning("Resume parse failed for %s: %s", filename, exc)
        raise ValueError(
            f"Couldn't read {filename!r}: the file looks corrupt or password-protected. "
            "Re-export it (File > Save as PDF/DOCX) or paste the text instead."
        ) from exc
    if not text.strip():
        raise ValueError(
            f"No text found in {filename!r}. It's probably a scanned image; ATS systems "
            "can't read those either. Export a text-based PDF or paste the text."
        )
    return text


def _extract_pdf(file_bytes: bytes) -> str:
    try:
        from pypdf import PdfReader  # lazy import
    except ImportError as exc:
        raise ValueError(
            "PDF support requires the 'pypdf' package. Install it or upload a DOCX/TXT resume."
        ) from exc

    reader = PdfReader(io.BytesIO(file_bytes))
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception as exc:  # noqa: BLE001
            logger.warning("PDF page extraction failed: %s", exc)
    return "\n".join(pages).strip()


def _extract_docx(file_bytes: bytes) -> str:
    try:
        import docx  # python-docx, lazy import
    except ImportError as exc:
        raise ValueError(
            "DOCX support requires 'python-docx'. Install it or upload a PDF/TXT resume."
        ) from exc

    document = docx.Document(io.BytesIO(file_bytes))
    # Body paragraphs alone miss tables (common in CV templates) and the
    # header/footer, which is where contact details usually live.
    parts: List[str] = []
    for section in document.sections:
        parts.extend(p.text for p in section.header.paragraphs)
    parts.extend(p.text for p in document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            cells = []
            for cell in row.cells:
                if cell.text not in cells:  # merged cells repeat their text
                    cells.append(cell.text)
            parts.append(" | ".join(c for c in cells if c.strip()))
    for section in document.sections:
        parts.extend(p.text for p in section.footer.paragraphs)
    return "\n".join(p for p in parts if p.strip()).strip()


# ---------------------------------------------------------------------------
# 2. Deterministic keyword match
# ---------------------------------------------------------------------------

_WORD_RE = re.compile(r"[A-Za-z][A-Za-z0-9+#./-]*")
_ICON_GLYPH_RE = re.compile("[\ue000-\uf8ff\ufffd]")


def _normalize_skill(skill: str) -> str:
    """Lowercase and strip surrounding noise like '(Expert level)'."""
    base = re.sub(r"\([^)]*\)", "", skill).strip().lower()
    base = re.sub(r"\s+", " ", base)
    return base


# Spelling variants: any member of a group satisfies a requirement for another.
_ALIAS_GROUPS = [
    {"kubernetes", "k8s"}, {"postgresql", "postgres"}, {"javascript", "js"},
    {"node.js", "nodejs", "node"}, {"react", "react.js", "reactjs"},
    {"vue", "vue.js", "vuejs"}, {"next.js", "nextjs"}, {"golang", "go"},
    {"machine learning", "ml"}, {"ci/cd", "cicd", "continuous integration"},
    {"aws", "amazon web services"}, {"gcp", "google cloud", "google cloud platform"},
    {"azure", "microsoft azure"}, {"scikit-learn", "sklearn"},
    {"rest", "restful", "rest api", "rest apis"},
    {"llm", "llms", "large language model", "large language models"},
    {"nlp", "natural language processing"}, {"tensorflow", "tf"},
]
_ALIASES: Dict[str, set] = {}
for _group in _ALIAS_GROUPS:
    for _term in _group:
        _ALIASES.setdefault(_term, set()).update(_group)

# Names that are also everyday words or letters ("go hiking", "R&D"): only a
# capitalized/upper-case mention in the original resume counts.
_CASE_SENSITIVE = {
    "go", "r", "c", "rust", "swift", "spring", "express", "dart", "elm", "ml",
    "js", "ai", "node", "rest", "tf", "chef", "puppet",
}

# Proficiency / hedging words inside a requirement, not part of the skill.
_QUALIFIER_RE = re.compile(
    r"\b(?:preferred|required|a plus|plus|nice to have|e\.g\.?|eg|i\.e\.?|etc\.?|such as"
    r"|like|including|expert|advanced|strong|solid|proficient|proficiency|basic"
    r"|intermediate|working knowledge of|knowledge of|experience with|experience in"
    r"|experience|familiarity with|level|years?|\d+\+?)\b",
    re.IGNORECASE,
)


_GENERIC_TAILS = {
    "api", "apis", "framework", "frameworks", "platform", "platforms", "tools",
    "tooling", "services", "database", "databases", "systems", "language",
    "languages", "development", "pipelines", "stack", "architecture",
}


# What's left after dropping the tail must still say something specific
# ("Data pipelines" -> "data" would match every resume).
_TOO_GENERIC = {
    "data", "web", "software", "programming", "modern", "scalable", "large",
    "high", "big", "new", "backend", "frontend", "mobile", "cloud", "cross",
}


def _skill_alternatives(skill: str) -> List[str]:
    """Every term that satisfies a requirement.

    "Containerization (Docker, Kubernetes)" -> containerization, docker,
    kubernetes; "AWS/GCP preferred" -> aws/gcp, aws, gcp; "Python (Expert
    level required)" -> python. The LLM writes requirements as
    "skill (detail)", and stripping the parenthetical alone used to discard the
    concrete tools the resume actually names.
    """
    inner = " , ".join(re.findall(r"\(([^)]*)\)", skill))
    base = re.sub(r"\([^)]*\)", " ", skill)
    out: List[str] = []
    for chunk in (base, inner):
        whole = _QUALIFIER_RE.sub(" ", chunk)
        pieces = [whole] + re.split(r"[,;/|&]|\bor\b|\band\b", whole, flags=re.IGNORECASE)
        for piece in pieces:
            term = re.sub(r"\s+", " ", piece).strip(" .:-").lower()
            if not term:
                continue
            words = term.split()
            candidates = [term]
            # "RESTful APIs" -> "restful"; "Relational databases" -> "relational"
            head = " ".join(words[:-1])
            if len(words) > 1 and words[-1] in _GENERIC_TAILS and head not in _TOO_GENERIC:
                candidates.append(head)
            candidates += [c[:-1] for c in list(candidates) if len(c) > 4 and c.endswith("s")]
            out.extend(c for c in candidates if c not in out)
    return out


def _term_present(term: str, resume_text: str, resume_text_lower: str) -> bool:
    boundary_l, boundary_r = r"(?<![A-Za-z0-9+#&])", r"(?![A-Za-z0-9+#&])"
    if term in _CASE_SENSITIVE:
        variants = {term.upper(), term.capitalize()}
        pattern = boundary_l + "(?:" + "|".join(map(re.escape, variants)) + ")" + boundary_r
        return re.search(pattern, resume_text) is not None
    pattern = boundary_l + re.escape(term) + boundary_r
    return re.search(pattern, resume_text_lower) is not None


def _skill_present(skill: str, resume_text_lower: str, resume_text: Optional[str] = None) -> bool:
    """Whole-token match for a skill (or any of its alternatives/aliases).

    Boundary class includes ``+`` and ``#`` so a search for "C" doesn't match
    inside "C#" / "C++", and a search for "C++" doesn't match inside a
    hypothetical "C+++". ``.`` is intentionally NOT a boundary char so normal
    sentence punctuation ("…in Python.") doesn't block a real match.
    """
    original = resume_text if resume_text is not None else resume_text_lower
    for term in _skill_alternatives(skill):
        for variant in _ALIASES.get(term, {term}):
            if _term_present(variant, original, resume_text_lower):
                return True
    return False


def ats_keyword_match(
    resume_text: str,
    required_skills: Iterable[str],
) -> Dict[str, object]:
    """Return matched / missing required skills and an overlap score 0-100."""
    resume_lower = resume_text.lower()
    required = [s for s in (required_skills or []) if s and isinstance(s, str)]

    matched: List[str] = []
    missing: List[str] = []
    for skill in required:
        (matched if _skill_present(skill, resume_lower, resume_text) else missing).append(skill)

    score = round(100 * len(matched) / len(required)) if required else 0
    return {
        "score": score,
        "matched": matched,
        "missing": missing,
        "required_count": len(required),
    }


# ---------------------------------------------------------------------------
# 3. ATS format checks
# ---------------------------------------------------------------------------

def ats_format_checks(resume_text: str) -> List[str]:
    """Return a list of ATS-unfriendly patterns detected in the resume text."""
    issues: List[str] = []
    if not resume_text.strip():
        return ["Resume text is empty — ATS will read nothing."]

    # Tabs / pipe-heavy lines often indicate tables that ATS parsers mangle.
    lines = resume_text.splitlines()
    pipey = sum(1 for line in lines if line.count("|") >= 3)
    if pipey >= 2:
        issues.append("Multiple pipe-separated lines detected — tables often fail in ATS parsers.")

    if any("\t\t" in line for line in lines):
        issues.append("Multi-tab columns detected — column layouts often fail in ATS parsers.")

    # Excessive ALL-CAPS body lines (headers are fine; many in a row aren't).
    caps_lines = [
        line for line in lines
        if len(line) > 20 and line == line.upper() and any(c.isalpha() for c in line)
    ]
    if len(caps_lines) >= 3:
        issues.append("Many ALL-CAPS body lines — consider sentence case for readability and ATS.")

    # Icon-font glyphs (Wingdings / Font Awesome bullets and contact icons) land
    # in the Unicode private-use area when PDFs/DOCX are converted to text.
    # (This used to test ``"" in resume_text`` — always true — so every resume
    # got this warning.)
    if _ICON_GLYPH_RE.search(resume_text):
        issues.append("Non-text glyphs detected — likely icons/decorations that ATS can't read.")

    word_count = len(_WORD_RE.findall(resume_text))
    if word_count < 150:
        issues.append(f"Resume is very short ({word_count} words) — ATS may rank it lower.")
    elif word_count > 1500:
        issues.append(f"Resume is very long ({word_count} words) — consider trimming.")

    return issues


# ---------------------------------------------------------------------------
# 4. Full resume analysis (deterministic score + LLM gap commentary)
# ---------------------------------------------------------------------------

def analyze_resume(
    resume_text: str,
    job_posting: str,
    required_skills: Iterable[str],
    model: str = "detailed",
) -> Dict[str, object]:
    """Produce the structured ATS analysis used by the report generator."""
    keyword_match = ats_keyword_match(resume_text, required_skills)
    format_issues = ats_format_checks(resume_text)

    # Cache the LLM commentary on (resume, posting, model) — the deterministic
    # parts are cheap and recomputed.
    cache_key = f"resume_gap_{hash(resume_text)}_{hash(job_posting)}_{model}"
    commentary: Optional[str] = cache.get(cache_key)
    if commentary is None:
        prompt = f"""
You are an ATS-optimization coach. Compare the candidate's resume against the
job posting and produce a concise markdown briefing.

Required skills already detected in the resume (deterministic): {keyword_match['matched']}
Required skills NOT detected in the resume (deterministic): {keyword_match['missing']}
Detected ATS formatting issues: {format_issues}

Job posting:
{wrap_untrusted(job_posting, "job_posting")}

Resume:
{wrap_untrusted(resume_text, "resume")}

Produce these sections (markdown):
1. **Overall fit** — one paragraph, grounded in the deterministic match above.
2. **Strengths to lead with** — 3-5 bullets the candidate should emphasize.
3. **Gaps and how to address them** — for each missing required skill, one
   concrete action (transferable experience to highlight, course to take,
   project to add). Do not invent skills the candidate doesn't have.
4. **Resume changes** — concrete edits (re-phrasings, sections to add/cut)
   that would improve the ATS score, including formatting fixes.

Rules:
- Do not fabricate facts about the candidate.
- Do not output `<think>` blocks or markdown code fences around the whole report.
"""
        commentary = get_completion(prompt, model)
        cache.set(cache_key, commentary)

    return {
        "ats_score": keyword_match["score"],
        "matched_skills": keyword_match["matched"],
        "missing_skills": keyword_match["missing"],
        "format_issues": format_issues,
        "commentary": commentary,
    }
